from flask import Flask, request, jsonify
from openpyxl import Workbook, load_workbook
import os

app = Flask(__name__)

# 엑셀 파일 생성 API
@app.route('/create_excel', methods=['POST'])
def create_excel_file():
    file_name = request.json.get('file_name', 'user_data.xlsx')

    if os.path.exists(file_name):
        return jsonify({"message": f"{file_name} 파일이 이미 존재합니다."}), 400

    # 새로운 엑셀 워크북 생성
    wb = Workbook()
    ws = wb.active

    # 헤더 추가
    ws.append(["닉네임", "아이디", "비밀번호", "최고 기록"])

    # 파일 저장
    wb.save(file_name)
    return jsonify({"message": f"{file_name} 파일을 생성했습니다."}), 201

# 사용자 정보 추가 API
@app.route('/add_user', methods=['POST'])
def add_user_to_excel():
    data = request.json
    file_name = data.get('file_name', 'user_data.xlsx')
    nickname = data.get('nickname')
    user_id = data.get('user_id')
    password = data.get('password')
    high_score = data.get('high_score')

    if not os.path.exists(file_name):
        return jsonify({"message": f"{file_name} 파일이 없습니다. 먼저 파일을 생성하세요."}), 404

    # 엑셀 파일 열기
    wb = load_workbook(file_name)
    ws = wb.active

    # 데이터 추가
    ws.append([nickname, user_id, password, high_score])

    # 파일 저장
    wb.save(file_name)
    return jsonify({
        "message": "사용자 정보가 추가되었습니다.",
        "user": {
            "nickname": nickname,
            "user_id": user_id,
            "high_score": high_score
        }
    }), 200

# 메인 함수
if __name__ == '__main__':
    app.run(debug=True)
